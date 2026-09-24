import tkinter as tk
from tkinter import ttk, messagebox
import openpyxl as op

window = tk.Tk()
window.title("Diet Monitoring System")
window.geometry("1280x720")
window.minsize(1050, 600)
window.configure(bg="#F4F6F8")

def display():
    workbook = op.load_workbook("Rosel_Database.xlsx")
    sheet = workbook.active

    for row in tree.get_children():
        tree.delete(row)

    for row in sheet.iter_rows(min_row=1, values_only=True):
        tree.insert("", tk.END, values=row)

def new_window(values=None):

    new_window = tk.Toplevel(window)
    new_window.title("Add Food Record")
    new_window.geometry("400x400")
    new_window.config(bg="white")
    new_window.resizable(False, False)

    title = tk.Label(new_window, text="🍎 Add Diet Record",font=("Arial", 16, "bold"),bg="white")
    title.grid(row=0, column=0, columnspan=2, pady=15)

    labels = ["Food Name","Meal Type","Serving Size","No. of Servings","Calories Per Serving","Total Calories","Date"]

    for i, text in enumerate(labels):

        label = tk.Label(new_window,text=text + ":",bg="white",font=("Arial", 10, "bold"))
        label.grid(row=i + 1,column=0,padx=20,pady=8,sticky="w")

    food_entry = tk.Entry(new_window, width=30)
    food_entry.grid(row=1, column=1)

    meal_type = ttk.Combobox(new_window, width=27,state="readonly",values=["Breakfast", "Lunch", "Dinner", "Snack"])
    meal_type.grid(row=2, column=1)
    meal_type.current(0)

    serving_size_entry = tk.Entry(new_window, width=30)
    serving_size_entry.grid(row=3, column=1)

    no_of_serving_entry = tk.Entry(new_window, width=30)
    no_of_serving_entry.grid(row=4, column=1)

    calories_entry = tk.Entry(new_window, width=30)
    calories_entry.grid(row=5, column=1)

    total_calories_entry = tk.Entry(new_window,width=30,state="readonly")
    total_calories_entry.grid(row=6, column=1)

    date_entry = tk.Entry(new_window, width=30)
    date_entry.grid(row=7, column=1)
    date_entry.insert(0, "06/00/2026")

    def calculate_total_calories(event=None):

        servings = no_of_serving_entry.get()
        calories = calories_entry.get()

        valid_servings = servings.replace(".", "", 1).isdigit()
        valid_calories = calories.replace(".", "", 1).isdigit()

        if servings != "" and calories != "":

            if valid_servings and valid_calories:

                total = float(servings) * float(calories)

                total_calories_entry.config(state="normal")
                total_calories_entry.delete(0, tk.END)
                total_calories_entry.insert(0, str(total))
                total_calories_entry.config(state="readonly")

            else:

                total_calories_entry.config(state="normal")
                total_calories_entry.delete(0, tk.END)
                total_calories_entry.config(state="readonly")

    no_of_serving_entry.bind("<KeyRelease>", calculate_total_calories)
    calories_entry.bind("<KeyRelease>", calculate_total_calories)

    def input_validation():
        calories = calories_entry.get()
        number_of_serving = no_of_serving_entry.get()

        if food_entry.get() == "":
            messagebox.showerror("Error","Food Name is required!")
            return False

        if serving_size_entry.get() == "":
            messagebox.showerror("Error","Serving Size is required!")
            return False

        if no_of_serving_entry.get() == "":
            messagebox.showerror("Error","No. of Servings is required!")
            return False
        elif not number_of_serving.isdigit():
            messagebox.showerror("Error","No. of Servings must be a number!")
            return False

        if calories_entry.get() == "":
            messagebox.showerror("Error","Calories is required!")
            return False
        elif not calories.isdigit():
            messagebox.showerror("Error","Calories must be a number!")
            return False

        return True

    def clear_entries(): 

        food_entry.delete(0, tk.END)
        meal_type.current(0)
        serving_size_entry.delete(0, tk.END)
        no_of_serving_entry.delete(0, tk.END)
        calories_entry.delete(0, tk.END)

        total_calories_entry.config(state="normal")
        total_calories_entry.delete(0, tk.END)
        total_calories_entry.config(state="readonly")

        date_entry.delete(0, tk.END)
        date_entry.insert(0, "06/00/2026")

    def save_record():
        if not input_validation():
            return

        food = food_entry.get()
        mealtype = meal_type.get()
        serving_size = serving_size_entry.get()
        no_servings = no_of_serving_entry.get()
        calories = calories_entry.get()
        total = total_calories_entry.get()
        date = date_entry.get()

        workbook = op.load_workbook("Rosel_Database.xlsx")
        sheet = workbook.active

        new_id = sheet.max_row

        sheet.append([new_id,food,mealtype,serving_size,no_servings,calories,total,date])
        workbook.save("Rosel_Database.xlsx")

        messagebox.showinfo("Success","Record added successfully!")
        display()
        new_window.destroy()

    def update_record():
        if not input_validation():
            return

        workbook = op.load_workbook("Rosel_Database.xlsx")
        sheet = workbook.active

        for row in sheet.iter_rows(min_row=1):
            if str(row[0].value) == str(values[0]):
                row[1].value = food_entry.get()
                row[2].value = meal_type.get()
                row[3].value = serving_size_entry.get()
                row[4].value = no_of_serving_entry.get()
                row[5].value = calories_entry.get()
                row[6].value = total_calories_entry.get()
                row[7].value = date_entry.get()
                break

        workbook.save("Rosel_Database.xlsx")

        messagebox.showinfo("Success","Record updated successfully!")
        display()
        new_window.destroy()

    if values:
        food_entry.insert(0, values[1])
        meal_type.set(values[2])
        serving_size_entry.insert(0, values[3])
        no_of_serving_entry.insert(0, values[4])
        calories_entry.insert(0, values[5])

        total_calories_entry.config(state="normal")
        total_calories_entry.insert(0, values[6])
        total_calories_entry.config(state="readonly")

        date_entry.delete(0, tk.END)
        date_entry.insert(0, values[7])

    button_frame = tk.Frame(new_window,bg="white")
    button_frame.grid(row=8,column=0,columnspan=2,pady=20)

    if values:
        save_btn = tk.Button(button_frame,text="Update",bg="#2196F3",fg="white",width=12,command=update_record)
    else:
        save_btn = tk.Button(button_frame,text="Save",bg="#4CAF50",fg="white",width=12,command=save_record)

    save_btn.grid(row=0, column=0, padx=10)

    clear_btn = tk.Button(button_frame,text="Clear",bg="#FF9800",fg="white",width=12,command=clear_entries)
    clear_btn.grid(row=0, column=1, padx=10)

    cancel_btn = tk.Button(button_frame,text="Cancel",bg="#F44336",fg="white",width=12,command=new_window.destroy)
    cancel_btn.grid(row=0, column=2, padx=10)


def update():
    selected = tree.focus()

    if not selected:
        messagebox.showerror("Error","Select a record first!")
        return
    values = tree.item(selected, "values")
    new_window(values)

def delete():
    selected = tree.focus()

    if not selected:
        messagebox.showerror("Error","Select a record first!")
        return

    values = tree.item(selected, "values")
    record_id = values[0]

    confirm = messagebox.askyesno("Confirm","Are you sure you want to delete this record?")

    if not confirm:
        return

    workbook = op.load_workbook("Rosel_Database.xlsx")
    sheet = workbook.active

    for i, row in enumerate(sheet.iter_rows(min_row=1),start=1):
        if str(row[0].value) == str(record_id):
            sheet.delete_rows(i)
            break

    workbook.save("Rosel_Database.xlsx")
    messagebox.showinfo("Success","Record deleted successfully!")
    display()

# =========================
# GUI DESIGN
# =========================

# Configure the main window grid
window.grid_rowconfigure(1, weight=1)
window.grid_columnconfigure(1, weight=1)

# Modern Treeview styling
style = ttk.Style()
style.theme_use("clam")

style.configure(
    "Treeview",
    background="white",
    foreground="#263238",
    rowheight=34,
    fieldbackground="white",
    font=("Arial", 10)
)

style.configure(
    "Treeview.Heading",
    background="#2E7D32",
    foreground="white",
    font=("Arial", 10, "bold"),
    padding=8
)

style.map(
    "Treeview",
    background=[("selected", "#C8E6C9")],
    foreground=[("selected", "#1B5E20")]
)

# Header
header_frame = tk.Frame(window, bg="#2E7D32", height=78)
header_frame.grid(row=0, column=0, columnspan=2, sticky="ew")
header_frame.grid_propagate(False)

title = tk.Label(
    header_frame,
    text="🥗  DIET MONITORING SYSTEM",
    font=("Arial", 22, "bold"),
    bg="#2E7D32",
    fg="white"
)
title.pack(expand=True)

# Sidebar
sidebar_frame = tk.Frame(window, bg="#E8F5E9", width=230)
sidebar_frame.grid(row=1, column=0, sticky="nsw")
sidebar_frame.grid_propagate(False)

sidebar_title = tk.Label(
    sidebar_frame,
    text="MENU",
    font=("Arial", 14, "bold"),
    bg="#E8F5E9",
    fg="#1B5E20"
)
sidebar_title.pack(pady=(28, 22))

button_style = {
    "font": ("Arial", 11, "bold"),
    "width": 18,
    "height": 2,
    "bd": 0,
    "relief": "flat",
    "cursor": "hand2",
    "activeforeground": "white"
}

add_btn = tk.Button(
    sidebar_frame,
    text="＋  ADD",
    bg="#43A047",
    fg="white",
    activebackground="#2E7D32",
    command=new_window,
    **button_style
)
add_btn.pack(fill="x", padx=22, pady=8)

update_btn = tk.Button(
    sidebar_frame,
    text="✎  UPDATE",
    bg="#1E88E5",
    fg="white",
    activebackground="#1565C0",
    command=update,
    **button_style
)
update_btn.pack(fill="x", padx=22, pady=8)

delete_btn = tk.Button(
    sidebar_frame,
    text="✕  DELETE",
    bg="#E53935",
    fg="white",
    activebackground="#B71C1C",
    command=delete,
    **button_style
)
delete_btn.pack(fill="x", padx=22, pady=8)

sidebar_note = tk.Label(
    sidebar_frame,
    text="Manage your daily food\nand calorie records",
    font=("Arial", 9),
    bg="#E8F5E9",
    fg="#546E7A",
    justify="center"
)
sidebar_note.pack(side="bottom", pady=25)

# Main content area
main_frame = tk.Frame(window, bg="#F4F6F8")
main_frame.grid(row=1, column=1, sticky="nsew", padx=(12, 18), pady=14)

main_frame.grid_rowconfigure(1, weight=1)
main_frame.grid_columnconfigure(0, weight=1)

table_title = tk.Label(
    main_frame,
    text="Food Records",
    font=("Arial", 16, "bold"),
    bg="#F4F6F8",
    fg="#263238",
    anchor="w"
)
table_title.grid(row=0, column=0, columnspan=2, sticky="ew", pady=(0, 10))

table_frame = tk.Frame(main_frame, bg="white", bd=1, relief="solid")
table_frame.grid(row=1, column=0, columnspan=2, sticky="nsew")

table_frame.grid_rowconfigure(0, weight=1)
table_frame.grid_columnconfigure(0, weight=1)

columns = (
    "ID", "Food Name", "Meal Type", "Serving Size",
    "No. of Servings", "Calories", "Total Calories", "Date"
)

tree = ttk.Treeview(
    table_frame,
    columns=columns,
    show="headings",
    selectmode="browse"
)

column_widths = {
    "ID": 55,
    "Food Name": 160,
    "Meal Type": 110,
    "Serving Size": 115,
    "No. of Servings": 125,
    "Calories": 95,
    "Total Calories": 125,
    "Date": 110
}

for heading in columns:
    tree.heading(heading, text=heading)
    tree.column(
        heading,
        width=column_widths[heading],
        minwidth=column_widths[heading],
        anchor="center"
    )

tree.grid(row=0, column=0, sticky="nsew")

vertical_scrollbar = ttk.Scrollbar(
    table_frame,
    orient="vertical",
    command=tree.yview
)
vertical_scrollbar.grid(row=0, column=1, sticky="ns")

horizontal_scrollbar = ttk.Scrollbar(
    table_frame,
    orient="horizontal",
    command=tree.xview
)
horizontal_scrollbar.grid(row=1, column=0, sticky="ew")

tree.configure(
    yscrollcommand=vertical_scrollbar.set,
    xscrollcommand=horizontal_scrollbar.set
)

display()
window.mainloop()
